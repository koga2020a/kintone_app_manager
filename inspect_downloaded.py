#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""取得済み output ディレクトリから Webhook と全文検索を行う。"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

import yaml

from kintone_dev_api import fetch_dev_app_json

SEARCH_SUFFIXES = {
    ".js",
    ".css",
    ".json",
    ".yaml",
    ".yml",
    ".txt",
    ".tsv",
    ".csv",
    ".html",
    ".md",
}

KIND_ORDER = ["JavaScript", "CSS", "設定JSON", "設定YAML", "その他テキスト"]

CATEGORY_LABELS = {
    "webhooks": "webhook",
    "webhook": "webhook",
    "javascript_info": "javascript",
    "field_codes_usage_at_javascript": "javascript",
    "customize": "customize",
    "form": "form",
    "form_fields": "form",
    "form_layout": "form",
    "views": "view",
    "graphs": "graph",
    "plugins": "plugin",
    "actions": "action",
    "settings": "settings",
    "app_acl": "app_acl",
    "record_acl": "record_acl",
    "field_acl": "field_acl",
    "app_notifications": "notification",
    "general_notifications": "notification",
    "record_notifications": "notification",
    "reminder_notifications": "notification",
    "process_management": "process",
}


def _app_id_from_output_name(name: str) -> Optional[str]:
    prefix = name.split("_", 1)[0]
    return prefix if prefix.isdigit() else None


def find_app_output_dir(output_dir: Path, app_id: str) -> Optional[Path]:
    """output 内の {app_id}_* ディレクトリのうち、更新が新しいものを返す。"""
    if not output_dir.exists():
        return None
    matches = [
        path
        for path in output_dir.iterdir()
        if path.is_dir() and path.name.startswith(f"{app_id}_")
    ]
    if not matches:
        return None
    return max(matches, key=lambda path: path.stat().st_mtime)


def list_app_output_dirs(output_dir: Path) -> List[Tuple[str, Path]]:
    """output 内の各アプリIDについて、最新フォルダを返す。"""
    if not output_dir.exists():
        return []
    by_id: Dict[str, List[Path]] = {}
    for path in output_dir.iterdir():
        if not path.is_dir():
            continue
        app_id = _app_id_from_output_name(path.name)
        if not app_id:
            continue
        by_id.setdefault(app_id, []).append(path)
    result: List[Tuple[str, Path]] = []
    for app_id in sorted(by_id, key=int):
        newest = max(by_id[app_id], key=lambda path: path.stat().st_mtime)
        result.append((app_id, newest))
    return result


def classify_file(app_dir: Path, file_path: Path) -> str:
    rel = file_path.relative_to(app_dir).as_posix()
    suffix = file_path.suffix.lower()
    if "javascript/" in rel or suffix == ".js":
        return "JavaScript"
    if suffix == ".css":
        return "CSS"
    if suffix in {".json"}:
        return "設定JSON"
    if suffix in {".yaml", ".yml"}:
        return "設定YAML"
    return "その他テキスト"


def classify_category(app_dir: Path, file_path: Path) -> str:
    """webhook / javascript など、設定内容の種別を返す。"""
    rel = file_path.relative_to(app_dir).as_posix().replace("\\", "/").lower()
    suffix = file_path.suffix.lower()
    if "/javascript/" in f"/{rel}" or suffix == ".js":
        return "javascript"
    if suffix == ".css":
        return "css"
    stem = file_path.stem.lower()
    match = re.match(r"^\d+_(.+)$", stem)
    if match:
        stem = match.group(1)
    return CATEGORY_LABELS.get(stem, stem or "other")


def load_json_or_yaml(path: Path) -> Any:
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() in {".yaml", ".yml"}:
        import yaml

        return yaml.safe_load(text)
    return json.loads(text)


def find_webhook_file(app_dir: Path, app_id: str) -> Optional[Path]:
    names = [
        f"{app_id}_webhooks.json",
        f"{app_id}_webhooks.yaml",
        f"{app_id}_webhook.json",
        f"{app_id}_webhook.yaml",
        "webhooks.json",
        "webhooks.yaml",
    ]
    folders = [app_dir / "json", app_dir]
    for folder in folders:
        if not folder.exists():
            continue
        for name in names:
            path = folder / name
            if path.exists():
                return path
    return None


def save_webhook_data(app_dir: Path, app_id: str, data: Any) -> Path:
    json_dir = app_dir / "json"
    json_dir.mkdir(parents=True, exist_ok=True)
    json_path = json_dir / f"{app_id}_webhooks.json"
    yaml_path = app_dir / f"{app_id}_webhooks.yaml"
    json_path.write_text(json.dumps(data, ensure_ascii=False, indent=4), encoding="utf-8")
    yaml_path.write_text(yaml.dump(data, allow_unicode=True), encoding="utf-8")
    return json_path


def fetch_webhooks_via_admin(
    subdomain: str,
    app_id: str,
    username: str,
    password: str,
) -> Tuple[Optional[Any], Optional[str]]:
    """管理画面と同じ内部APIで Webhook 一覧を取る。"""
    return fetch_dev_app_json(subdomain, username, password, app_id, "webhook/list.json")


def fetch_and_save_webhooks(
    app_dir: Path,
    app_id: str,
    subdomain: str,
    api_token: Optional[str] = None,
    username: Optional[str] = None,
    password: Optional[str] = None,
) -> Tuple[Optional[Path], Optional[Any], Optional[str]]:
    """Webhook一覧を取得して output に保存する。

    api_token は互換のため残しているが、この内部APIでは使用しない
    (API トークン認証ではアクセスできない)。
    """
    if not username or not password:
        return (
            None,
            None,
            "Webhook一覧の取得には .kintone.env の username / password が必要です",
        )

    data, error = fetch_webhooks_via_admin(subdomain, app_id, username, password)
    if data is not None:
        saved = save_webhook_data(app_dir, app_id, data)
        return saved, data, None
    return None, None, error or "Webhook一覧を取得できませんでした"


def _person_name(person: Any) -> str:
    if isinstance(person, dict):
        return str(person.get("name") or person.get("code") or "")
    if person is None:
        return ""
    return str(person)


def extract_webhook_rows(data: Any) -> List[Dict[str, Any]]:
    """内部API のレスポンス（result 部）から表示用の行を作る。"""
    if data is None:
        return []
    if isinstance(data, dict):
        result = data.get("result")
        if isinstance(result, (dict, list)) and not data.get("webhooks"):
            return extract_webhook_rows(result)
        items = data.get("webhooks") or data.get("webhook") or []
        if isinstance(items, dict):
            items = [items]
        if not items and any(key in data for key in ("url", "localId", "webhookId", "id")):
            items = [data]
    elif isinstance(data, list):
        items = data
    else:
        return []

    rows = []
    for item in items:
        if not isinstance(item, dict):
            continue
        events = item.get("types") or item.get("events") or item.get("event") or []
        if isinstance(events, list):
            events_text = ", ".join(str(ev) for ev in events)
        else:
            events_text = str(events)
        headers = item.get("headers") or []
        if isinstance(headers, list):
            header_text = ", ".join(
                f"{h.get('name', '')}={h.get('value', '')}" if isinstance(h, dict) else str(h)
                for h in headers
            )
        else:
            header_text = str(headers)
        rows.append(
            {
                "id": item.get("localId") or item.get("webhookId") or item.get("id") or "",
                "name": item.get("description") or item.get("name") or "",
                "url": item.get("url") or "",
                "events": events_text,
                "enabled": item.get("enabled", ""),
                "creator": _person_name(item.get("creator")),
                "modifier": _person_name(item.get("modifier")),
                "created_at": item.get("createdAt") or "",
                "modified_at": item.get("modifiedAt") or "",
                "headers": header_text,
            }
        )
    return rows


def is_stale_webhook_file(data: Any) -> bool:
    """旧方式（管理画面HTMLの走査）で保存されたデータかどうか。"""
    if isinstance(data, dict):
        source = data.get("source")
        if isinstance(source, str) and "/k/admin/app/webhook" in source:
            return True
        candidates = [data]
        result = data.get("result")
        if isinstance(result, dict):
            candidates.append(result)
        for candidate in candidates:
            items = candidate.get("webhooks") or candidate.get("webhook") or []
            if isinstance(items, dict):
                items = [items]
            if isinstance(items, list):
                for item in items:
                    if isinstance(item, dict) and item.get("source") == "admin_html":
                        return True
    elif isinstance(data, list):
        for item in data:
            if isinstance(item, dict) and item.get("source") == "admin_html":
                return True
    return False


def _yaml_twin_exists(file_path: Path, app_dir: Path) -> bool:
    """同じ設定の YAML がある JSON は重複なので除外する。"""
    if file_path.suffix.lower() != ".json":
        return False
    stem = file_path.stem
    candidates = [
        file_path.with_suffix(".yaml"),
        file_path.with_suffix(".yml"),
        app_dir / f"{stem}.yaml",
        app_dir / f"{stem}.yml",
    ]
    return any(path.is_file() for path in candidates)


def search_downloaded(
    app_dir: Path, keywords: Iterable[str]
) -> List[Dict[str, Any]]:
    """キーワードのいずれかにヒットした行を返す。"""
    cleaned = [kw for kw in (k.strip() for k in keywords) if kw]
    if not cleaned:
        return []

    hits: List[Dict[str, Any]] = []
    for file_path in sorted(app_dir.rglob("*")):
        if not file_path.is_file():
            continue
        if file_path.suffix.lower() not in SEARCH_SUFFIXES:
            continue
        if _yaml_twin_exists(file_path, app_dir):
            continue
        try:
            lines = file_path.read_text(encoding="utf-8").splitlines()
        except (UnicodeDecodeError, OSError):
            continue
        kind = classify_file(app_dir, file_path)
        category = classify_category(app_dir, file_path)
        rel = file_path.relative_to(app_dir).as_posix()
        for line_no, line in enumerate(lines, start=1):
            matched = [kw for kw in cleaned if kw.lower() in line.lower()]
            if not matched:
                continue
            hits.append(
                {
                    "kind": kind,
                    "category": category,
                    "file": rel,
                    "line": line_no,
                    "text": line.strip(),
                    "keywords": matched,
                }
            )
    return hits


def summarize_hits(hits: List[Dict[str, Any]]) -> List[Tuple[str, int, int]]:
    """ファイル形式ごとのファイル数とヒット行数。"""
    files_by_kind: Dict[str, set] = {kind: set() for kind in KIND_ORDER}
    lines_by_kind: Dict[str, int] = {kind: 0 for kind in KIND_ORDER}
    for hit in hits:
        kind = hit["kind"]
        files_by_kind.setdefault(kind, set()).add(hit["file"])
        lines_by_kind[kind] = lines_by_kind.get(kind, 0) + 1
    summary = []
    for kind in KIND_ORDER:
        if lines_by_kind.get(kind):
            summary.append((kind, len(files_by_kind[kind]), lines_by_kind[kind]))
    return summary


def summarize_categories(hits: List[Dict[str, Any]]) -> List[Tuple[str, int, int]]:
    """webhook / javascript など種別ごとのファイル数とヒット行数。"""
    files_by_cat: Dict[str, set] = {}
    lines_by_cat: Dict[str, int] = {}
    for hit in hits:
        category = hit.get("category") or hit.get("kind") or "other"
        files_by_cat.setdefault(category, set()).add((hit.get("app_dir") or "", hit.get("file") or ""))
        lines_by_cat[category] = lines_by_cat.get(category, 0) + 1
    return [
        (category, len(files_by_cat[category]), lines_by_cat[category])
        for category in sorted(files_by_cat)
    ]


def hits_for_keyword(hits: List[Dict[str, Any]], keyword: str) -> List[Dict[str, Any]]:
    return [hit for hit in hits if keyword in (hit.get("keywords") or [])]


def summarize_keyword_hits(keyword: str, hits: List[Dict[str, Any]]) -> Dict[str, Any]:
    subset = hits_for_keyword(hits, keyword) if keyword != "（全体）" else hits
    apps = sorted({str(hit.get("app_id") or "") for hit in subset if hit.get("app_id")})
    files = {(hit.get("app_dir") or "", hit.get("file") or "") for hit in subset}
    kinds = {hit.get("kind") or "" for hit in subset if hit.get("kind")}
    categories = {hit.get("category") or "" for hit in subset if hit.get("category")}
    kind_text = "、".join(
        f"{kind}: {file_count}ファイル/{line_count}行"
        for kind, file_count, line_count in summarize_hits(subset)
    )
    category_text = "、".join(
        f"{category}: {file_count}ファイル/{line_count}行"
        for category, file_count, line_count in summarize_categories(subset)
    )
    return {
        "keyword": keyword,
        "hit_count": len(subset),
        "app_count": len(apps),
        "kind_count": len(kinds),
        "category_count": len(categories),
        "file_count": len(files),
        "kind_detail": kind_text,
        "category_detail": category_text,
        "app_ids": ", ".join(apps),
    }


def _safe_sheet_name(name: str, used: Set[str]) -> str:
    cleaned = re.sub(r'[:\\/?*\[\]]', "_", (name or "").strip()) or "検索語"
    cleaned = cleaned[:31]
    candidate = cleaned
    index = 2
    while candidate in used:
        suffix = f"_{index}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        index += 1
    used.add(candidate)
    return candidate


def _app_label(hit: Dict[str, Any]) -> str:
    app_dir = str(hit.get("app_dir") or "")
    app_id = str(hit.get("app_id") or "")
    prefix = f"{app_id}_"
    if app_id and app_dir.startswith(prefix):
        return app_dir[len(prefix):]
    return app_dir


def _app_sort_key(app_id: str, app_name: str = "") -> Tuple[int, Any, str]:
    text = str(app_id or "")
    if text.isdigit():
        return (0, int(text), app_name)
    return (1, text, app_name)


def summarize_app_hits(keyword: str, hits: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """検索語 × アプリ × 種別の合致情報。"""
    subset = hits_for_keyword(hits, keyword) if keyword != "（全体）" else hits
    grouped: Dict[Tuple[str, str, str], List[Dict[str, Any]]] = {}
    for hit in subset:
        key = (
            str(hit.get("app_id") or ""),
            _app_label(hit),
            str(hit.get("category") or ""),
        )
        grouped.setdefault(key, []).append(hit)

    rows: List[Dict[str, Any]] = []
    for (app_id, app_name, category), group in sorted(
        grouped.items(),
        key=lambda item: (*_app_sort_key(item[0][0], item[0][1]), item[0][2]),
    ):
        files = {(hit.get("app_dir") or "", hit.get("file") or "") for hit in group}
        kind_text = "、".join(
            f"{kind}: {file_count}ファイル/{line_count}行"
            for kind, file_count, line_count in summarize_hits(group)
        )
        rows.append(
            {
                "keyword": keyword,
                "app_id": app_id,
                "app_name": app_name,
                "category": category,
                "hit_count": len(group),
                "file_count": len(files),
                "kind_detail": kind_text,
            }
        )
    return rows


def export_search_hits_to_excel(
    keywords: Iterable[str],
    hits: List[Dict[str, Any]],
    output_path: Path,
) -> Path:
    """検索語ごとのサマリとヒット一覧を Excel に出力する。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    words = [kw for kw in (k.strip() for k in keywords) if kw]
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    header_fill = PatternFill(start_color="E6F3FF", end_color="E6F3FF", fill_type="solid")
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    app_fills = [
        PatternFill(start_color="DEEBF7", end_color="DEEBF7", fill_type="solid"),
        PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    ]
    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_align = Alignment(vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def fills_by_app(keys: List[Any]) -> List[PatternFill]:
        fills: List[PatternFill] = []
        last = object()
        index = -1
        for key in keys:
            if key != last:
                index += 1
                last = key
            fills.append(app_fills[index % 2])
        return fills

    def write_table(
        ws,
        headers: List[str],
        rows: List[List[Any]],
        widths: List[int],
        start_row: int = 1,
        row_fills: Optional[List[PatternFill]] = None,
        auto_filter: bool = True,
        freeze: bool = True,
    ) -> int:
        header_row = start_row
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin
        for offset, row in enumerate(rows):
            row_idx = header_row + 1 + offset
            fill = row_fills[offset] if row_fills and offset < len(row_fills) else None
            for col, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.alignment = data_align
                cell.border = thin
                if fill is not None:
                    cell.fill = fill
        for col, width in enumerate(widths, 1):
            current = ws.column_dimensions[get_column_letter(col)].width
            if current is None or current < width:
                ws.column_dimensions[get_column_letter(col)].width = width
        last_row = header_row + len(rows)
        if auto_filter and rows:
            ws.auto_filter.ref = (
                f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
            )
        if freeze:
            ws.freeze_panes = f"A{header_row + 1}"
        ws.row_dimensions[header_row].height = 22
        return last_row

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "サマリ"
    keyword_headers = [
        "検索語",
        "ヒット行数",
        "アプリ数",
        "種別数",
        "種別内訳",
        "ファイル種類数",
        "ファイル数",
        "ファイル種類内訳",
        "対象アプリ",
    ]
    keyword_rows: List[List[Any]] = []
    for word in words:
        row = summarize_keyword_hits(word, hits)
        keyword_rows.append(
            [
                row["keyword"],
                row["hit_count"],
                row["app_count"],
                row["category_count"],
                row["category_detail"],
                row["kind_count"],
                row["file_count"],
                row["kind_detail"],
                row["app_ids"],
            ]
        )
    total = summarize_keyword_hits("（全体）", hits)
    keyword_rows.append(
        [
            total["keyword"],
            total["hit_count"],
            total["app_count"],
            total["category_count"],
            total["category_detail"],
            total["kind_count"],
            total["file_count"],
            total["kind_detail"],
            total["app_ids"],
        ]
    )
    write_table(
        summary_ws,
        keyword_headers,
        keyword_rows,
        [24, 12, 10, 10, 40, 14, 12, 40, 24],
        auto_filter=False,
        freeze=False,
    )

    app_headers = [
        "検索語",
        "アプリID",
        "アプリ名",
        "種別",
        "ヒット行数",
        "ファイル数",
        "ファイル種類内訳",
    ]
    app_rows: List[List[Any]] = []
    app_keys: List[str] = []
    for word in [*words, "（全体）"]:
        for row in summarize_app_hits(word, hits):
            app_rows.append(
                [
                    row["keyword"],
                    row["app_id"],
                    row["app_name"],
                    row["category"],
                    row["hit_count"],
                    row["file_count"],
                    row["kind_detail"],
                ]
            )
            app_keys.append(str(row["app_id"]))

    title_row = len(keyword_rows) + 3
    title_cell = summary_ws.cell(row=title_row, column=1, value="アプリ毎の合致")
    title_cell.font = header_font
    title_cell.fill = section_fill
    last_col = get_column_letter(len(app_headers))
    summary_ws.merge_cells(f"A{title_row}:{last_col}{title_row}")
    for col in range(1, len(app_headers) + 1):
        cell = summary_ws.cell(row=title_row, column=col)
        cell.fill = section_fill
        cell.border = thin

    write_table(
        summary_ws,
        app_headers,
        app_rows,
        [24, 12, 24, 16, 12, 12, 40],
        start_row=title_row + 1,
        row_fills=fills_by_app(app_keys),
        auto_filter=True,
        freeze=True,
    )

    used_names = {summary_ws.title}
    hit_headers = ["アプリID", "アプリ名", "種別", "種類", "ファイル", "行", "内容", "検索語"]
    for word in words:
        ws = wb.create_sheet(_safe_sheet_name(word, used_names))
        subset = hits_for_keyword(hits, word)
        hit_rows = [
            [
                hit.get("app_id") or "",
                _app_label(hit),
                hit.get("category") or "",
                hit.get("kind") or "",
                hit.get("file") or "",
                hit.get("line") or "",
                hit.get("text") or "",
                word,
            ]
            for hit in subset
        ]
        write_table(
            ws,
            hit_headers,
            hit_rows,
            [12, 24, 16, 14, 40, 8, 80, 20],
            row_fills=fills_by_app([str(hit.get("app_id") or "") for hit in subset]),
        )

    wb.save(output_path)
    return output_path
