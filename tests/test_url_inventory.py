#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""url_inventory の単体テスト。

    python3 -m unittest tests.test_url_inventory

Excel 出力のテストは openpyxl が無い環境では skip する。
"""

from __future__ import annotations

import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import inspect_downloaded  # noqa: E402
import url_inventory as ui  # noqa: E402

try:
    import openpyxl  # noqa: F401

    HAS_OPENPYXL = True
except Exception:  # pragma: no cover - 環境依存
    HAS_OPENPYXL = False


APP_JS = """const API_BASE = 'https://old.example.com';
const LINK = 'https://docs.example.com/help';
var ns = 'http://www.w3.org/2000/svg';
// https://ignored.example.com
fetch('https://old.example.com/api/a');
fetch(API_BASE + '/b', {method: 'POST'});
kintone.proxy('https://api.old.example.com/x', 'PUT', {}, {});
kintone.api(kintone.api.url('/k/v1/record', true), 'GET', {});
var xhr = new XMLHttpRequest();
xhr.open('DELETE', 'https://new.example.com/y');
$.ajax({url: 'https://old.example.com/z', type: 'post'});
axios.get(`https://${host}/dyn`);
window.open('https://old.example.com/page');
"""

CUSTOMIZE_YAML = """desktop:
  css: []
  js:
  - type: URL
    url: https://cdn.old.example.com/lib.js
  - file:
      contentType: text/javascript
      fileKey: 2026081901014214409C8A8BEF4A978ECFE35C38ACAC0F048
      name: test.js
      size: '13'
    type: FILE
mobile:
  css:
  - type: URL
    url: https://cdn.old.example.com/style.css
  js: []
revision: '10'
scope: ALL
"""

VIEWS_YAML = """revision: '10'
views:
  カスタム一覧:
    html: '<div><script src="https://cdn.example.com/x.js"></script></div>'
    id: '5519897'
    index: '0'
    name: カスタム一覧
    pager: true
    type: CUSTOM
  顧客一覧:
    fields:
    - record_no
    id: '2189'
    index: '1'
    name: 顧客一覧
    sort: 作成日時 desc
    type: LIST
"""

WEBHOOKS_YAML = """webhooks:
- createdAt: '2026-08-18T22:28:50.000Z'
  creator:
    code: tester@example.com
    name: テスト太郎
  description: 旧サーバ通知
  enabled: true
  localId: '1001'
  modifiedAt: '2026-08-18T22:28:50.000Z'
  modifier:
    code: tester@example.com
    name: テスト太郎
  types:
  - ADD_RECORD
  - UPDATE_RECORD
  url: https://old.example.com/hook
- createdAt: '2026-08-18T22:29:50.000Z'
  creator:
    code: tester@example.com
    name: テスト太郎
  description: 別サーバ通知
  enabled: false
  localId: '1002'
  modifiedAt: '2026-08-18T22:29:50.000Z'
  modifier:
    code: tester@example.com
    name: テスト太郎
  types:
  - DELETE_RECORD
  url: https://other.example.org/hook
"""

DOMAIN_MAP = {"old.example.com": "new.example.com"}


class UrlInventoryTestBase(unittest.TestCase):
    """テスト用の output ディレクトリを 1 度だけ作る。"""

    @classmethod
    def setUpClass(cls) -> None:
        cls.tmp = Path(tempfile.mkdtemp(prefix="url_inventory_test_"))
        cls.output = cls.tmp / "output"
        app = cls.output / "123_テストアプリ"
        (app / "javascript").mkdir(parents=True)
        (app / "javascript" / "app.js").write_text(APP_JS, encoding="utf-8")
        (app / "javascript" / "vendor.min.js").write_text(
            'var _min="' + "x" * 2500 + '";', encoding="utf-8"
        )
        (app / "123_customize.yaml").write_text(CUSTOMIZE_YAML, encoding="utf-8")
        (app / "123_views.yaml").write_text(VIEWS_YAML, encoding="utf-8")
        (app / "123_webhooks.yaml").write_text(WEBHOOKS_YAML, encoding="utf-8")

        # Webhook を取得していないアプリ
        (cls.output / "456_未取得アプリ" / "javascript").mkdir(parents=True)

        cls.targets = inspect_downloaded.list_app_output_dirs(cls.output)
        cls.inv = ui.build_inventory(cls.targets, DOMAIN_MAP)
        cls.inv_nomap = ui.build_inventory(cls.targets)

    @classmethod
    def tearDownClass(cls) -> None:
        shutil.rmtree(cls.tmp, ignore_errors=True)

    # -- ヘルパー ---------------------------------------------------------
    def js_rows(self):
        return [r for r in self.inv.rows if r.category == ui.CAT_JS]

    def js_row(self, line: int, kind: str = ""):
        for row in self.js_rows():
            if row.line == line and (not kind or row.kind == kind):
                return row
        self.fail(f"JS呼び出しの行 {line} (kind={kind or '任意'}) が見つかりません")

    def ref_rows(self):
        return [r for r in self.inv.rows if r.category == ui.CAT_REF]

    def webhook_rows(self):
        return [r for r in self.inv.rows if r.category == ui.CAT_WEBHOOK]


class TestTargets(UrlInventoryTestBase):
    def test_targets(self):
        self.assertEqual([app_id for app_id, _ in self.targets], ["123", "456"])

    def test_apps(self):
        self.assertEqual(self.inv.apps, [("123", "テストアプリ"), ("456", "未取得アプリ")])

    def test_generated_at(self):
        self.assertIsNotNone(self.inv.generated_at)


class TestJavaScriptScan(UrlInventoryTestBase):
    def test_row_count(self):
        # 呼び出し 8 件 + URL文字列 2 件（コメント行と w3.org は除外）
        self.assertEqual(len(self.js_rows()), 10)

    def test_all_from_app_js(self):
        for row in self.js_rows():
            self.assertEqual(row.file, "javascript/app.js")
            self.assertEqual(row.app_id, "123")
            self.assertEqual(row.app_name, "テストアプリ")

    def test_fetch_literal(self):
        row = self.js_row(5)
        self.assertEqual(row.kind, "fetch")
        self.assertEqual(row.method, "GET")
        self.assertEqual(row.url, "https://old.example.com/api/a")
        self.assertEqual(row.resolved_url, "")
        self.assertEqual(row.host, "old.example.com")
        self.assertIn("fetch(", row.code)

    def test_fetch_variable_resolved(self):
        row = self.js_row(6)
        self.assertEqual(row.kind, "fetch")
        self.assertEqual(row.method, "POST")
        self.assertEqual(row.url, "API_BASE + '/b'")
        self.assertEqual(row.resolved_url, "https://old.example.com/b")
        self.assertEqual(row.host, "old.example.com")

    def test_kintone_proxy(self):
        row = self.js_row(7)
        self.assertEqual(row.kind, "kintone.proxy")
        self.assertEqual(row.method, "PUT")
        self.assertEqual(row.url, "https://api.old.example.com/x")
        self.assertEqual(row.host, "api.old.example.com")

    def test_kintone_api(self):
        row = self.js_row(8)
        self.assertEqual(row.kind, "kintone.api")
        self.assertEqual(row.method, "GET")
        self.assertIn("kintone.api.url(", row.url)
        self.assertEqual(row.host, "(kintone)")

    def test_xhr_open(self):
        row = self.js_row(10)
        self.assertEqual(row.kind, "XMLHttpRequest")
        self.assertEqual(row.method, "DELETE")
        self.assertEqual(row.url, "https://new.example.com/y")
        self.assertEqual(row.host, "new.example.com")

    def test_jquery_ajax(self):
        row = self.js_row(11)
        self.assertEqual(row.kind, "jQuery.ajax")
        self.assertEqual(row.method, "POST")
        self.assertEqual(row.url, "https://old.example.com/z")
        self.assertEqual(row.host, "old.example.com")

    def test_axios_dynamic_host(self):
        row = self.js_row(12)
        self.assertEqual(row.kind, "axios")
        self.assertEqual(row.method, "GET")
        self.assertIn("${host}", row.url)
        self.assertEqual(row.host, "(動的)")
        self.assertEqual(row.status, ui.STATUS_NONE)

    def test_window_open(self):
        row = self.js_row(13)
        self.assertEqual(row.kind, "window.open")
        self.assertEqual(row.method, "-")
        self.assertEqual(row.url, "https://old.example.com/page")

    def test_url_literal(self):
        row = self.js_row(2)
        self.assertEqual(row.kind, "URL文字列")
        self.assertEqual(row.method, "-")
        self.assertEqual(row.url, "https://docs.example.com/help")
        self.assertEqual(row.host, "docs.example.com")
        self.assertEqual(row.status, ui.STATUS_NONE)

    def test_comment_line_ignored(self):
        for row in self.inv.rows:
            self.assertNotIn("ignored.example.com", row.url)

    def test_namespace_host_ignored(self):
        for row in self.inv.rows:
            self.assertNotEqual(row.host, "www.w3.org")

    def test_call_line_not_duplicated_as_literal(self):
        # 呼び出しで拾った行は URL文字列 として二重に出さない
        kinds = [row.kind for row in self.js_rows() if row.line == 5]
        self.assertEqual(kinds, ["fetch"])

    def test_minified_file_skipped(self):
        skipped = [item for item in self.inv.skipped_files if "vendor.min.js" in item[1]]
        self.assertEqual(len(skipped), 1)
        self.assertEqual(skipped[0][0], "123")
        self.assertEqual(skipped[0][1], "javascript/vendor.min.js")
        self.assertIn("ミニファイ", skipped[0][2])


class TestWebhooks(UrlInventoryTestBase):
    def test_rows(self):
        rows = self.webhook_rows()
        self.assertEqual(len(rows), 2)
        urls = sorted(row.url for row in rows)
        self.assertEqual(
            urls, ["https://old.example.com/hook", "https://other.example.org/hook"]
        )

    def test_fields(self):
        row = [r for r in self.webhook_rows() if r.url.endswith("/hook") and r.host == "old.example.com"][0]
        self.assertEqual(row.kind, "Webhook")
        self.assertEqual(row.file, "")
        self.assertEqual(row.line, 0)
        self.assertEqual(row.method, "-")
        self.assertEqual(row.extra["id"], "1001")
        self.assertEqual(row.extra["name"], "旧サーバ通知")
        self.assertEqual(row.extra["enabled"], "有効")
        self.assertIn("ADD_RECORD", row.extra["events"])
        self.assertEqual(row.extra["modifier"], "テスト太郎")
        self.assertTrue(row.extra["modified_at"])
        self.assertEqual(row.code, row.extra["events"])

    def test_disabled(self):
        row = [r for r in self.webhook_rows() if r.host == "other.example.org"][0]
        self.assertEqual(row.extra["enabled"], "無効")
        self.assertEqual(row.status, ui.STATUS_NONE)

    def test_missing_webhooks(self):
        self.assertEqual(self.inv.missing_webhooks, ["456"])


class TestExternalRefs(UrlInventoryTestBase):
    def test_row_count(self):
        self.assertEqual(len(self.ref_rows()), 3)

    def test_customize_js(self):
        rows = [r for r in self.ref_rows() if r.kind == "カスタマイズJS"]
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row.url, "https://cdn.old.example.com/lib.js")
        self.assertEqual(row.host, "cdn.old.example.com")
        self.assertEqual(row.file, "123_customize.yaml")
        self.assertEqual(row.extra["platform"], "desktop")
        self.assertEqual(row.extra["file_type"], "js")
        self.assertGreater(row.line, 0)
        self.assertEqual(row.status, ui.STATUS_TODO)
        self.assertEqual(row.to_be_url, "https://cdn.new.example.com/lib.js")

    def test_customize_css_mobile(self):
        rows = [r for r in self.ref_rows() if r.kind == "カスタマイズCSS"]
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row.url, "https://cdn.old.example.com/style.css")
        self.assertEqual(row.extra["platform"], "mobile")
        self.assertEqual(row.extra["file_type"], "css")

    def test_view_html(self):
        rows = [r for r in self.ref_rows() if r.kind == "ビューHTML"]
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row.url, "https://cdn.example.com/x.js")
        self.assertEqual(row.host, "cdn.example.com")
        self.assertEqual(row.file, "123_views.yaml")
        self.assertEqual(row.extra["platform"], "カスタム一覧")
        self.assertEqual(row.extra["file_type"], "html")
        self.assertEqual(row.status, ui.STATUS_NONE)

    def test_file_type_url_only(self):
        # type: FILE のカスタマイズJS は拾わない
        for row in self.ref_rows():
            self.assertNotIn("test.js", row.url)


class TestDomainMap(UrlInventoryTestBase):
    def test_exact_host(self):
        row = self.js_row(5)
        self.assertEqual(row.status, ui.STATUS_TODO)
        self.assertEqual(row.to_be_url, "https://new.example.com/api/a")

    def test_resolved_url_is_replaced(self):
        row = self.js_row(6)
        self.assertEqual(row.status, ui.STATUS_TODO)
        self.assertEqual(row.to_be_url, "https://new.example.com/b")

    def test_subdomain(self):
        row = self.js_row(7)
        self.assertEqual(row.status, ui.STATUS_TODO)
        self.assertEqual(row.to_be_url, "https://api.new.example.com/x")

    def test_already_done(self):
        row = self.js_row(10)
        self.assertEqual(row.status, ui.STATUS_DONE)
        self.assertEqual(row.to_be_url, "")

    def test_out_of_scope(self):
        row = [r for r in self.webhook_rows() if r.host == "other.example.org"][0]
        self.assertEqual(row.status, ui.STATUS_NONE)
        self.assertEqual(row.to_be_url, "")

    def test_webhook_to_be_url(self):
        row = [r for r in self.webhook_rows() if r.host == "old.example.com"][0]
        self.assertEqual(row.status, ui.STATUS_TODO)
        self.assertEqual(row.to_be_url, "https://new.example.com/hook")

    def test_no_domain_map(self):
        self.assertEqual(self.inv_nomap.domain_map, {})
        for row in self.inv_nomap.rows:
            self.assertEqual(row.status, "")
            self.assertEqual(row.to_be_url, "")


class TestNormalizeDomainMap(unittest.TestCase):
    def test_scheme_and_slash(self):
        self.assertEqual(
            ui.normalize_domain_map({"https://Old.Example.com/": "new.example.com"}),
            {"old.example.com": "new.example.com"},
        )

    def test_none(self):
        self.assertEqual(ui.normalize_domain_map(None), {})

    def test_not_dict(self):
        self.assertEqual(ui.normalize_domain_map("old.example.com"), {})
        self.assertEqual(ui.normalize_domain_map([1, 2]), {})

    def test_port_and_path(self):
        self.assertEqual(
            ui.normalize_domain_map({"http://a.example.com:8080/path": "B.Example.NET/"}),
            {"a.example.com": "b.example.net"},
        )

    def test_empty_value_dropped(self):
        self.assertEqual(ui.normalize_domain_map({"a.example.com": ""}), {})


class TestSummarizeHosts(UrlInventoryTestBase):
    def test_order_and_counts(self):
        hosts = ui.summarize_hosts(self.inv)
        self.assertTrue(hosts)
        first = hosts[0]
        self.assertEqual(first["host"], "old.example.com")
        self.assertEqual(first["status"], ui.STATUS_TODO)
        self.assertEqual(first["to_be_host"], "new.example.com")
        self.assertEqual(first["webhook"], 1)
        self.assertEqual(first["js"], 5)
        self.assertEqual(first["ref"], 0)
        self.assertEqual(first["total"], 6)
        self.assertEqual(first["apps"], ["123"])

    def test_status_grouping(self):
        hosts = ui.summarize_hosts(self.inv)
        order = [ui.STATUS_ORDER.get(str(entry["status"]), 3) for entry in hosts]
        self.assertEqual(order, sorted(order))
        todo = [entry["host"] for entry in hosts if entry["status"] == ui.STATUS_TODO]
        self.assertEqual(
            todo, ["old.example.com", "cdn.old.example.com", "api.old.example.com"]
        )
        done = [entry["host"] for entry in hosts if entry["status"] == ui.STATUS_DONE]
        self.assertEqual(done, ["new.example.com"])

    def test_totals_match_rows(self):
        hosts = ui.summarize_hosts(self.inv)
        self.assertEqual(sum(int(entry["total"]) for entry in hosts), len(self.inv.rows))

    def test_without_domain_map(self):
        hosts = ui.summarize_hosts(self.inv_nomap)
        for entry in hosts:
            self.assertEqual(entry["status"], "")
            self.assertEqual(entry["to_be_host"], "")


class TestFormatSummary(UrlInventoryTestBase):
    def test_lines(self):
        lines = ui.format_inventory_summary(self.inv)
        self.assertTrue(lines)
        for line in lines:
            self.assertIsInstance(line, str)
        text = "\n".join(lines)
        self.assertIn("old.example.com", text)
        self.assertIn(ui.STATUS_TODO, text)
        self.assertIn("vendor.min.js", text)
        self.assertIn("456", text)

    def test_lines_without_domain_map(self):
        lines = ui.format_inventory_summary(self.inv_nomap)
        self.assertTrue(lines)


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl がインストールされていません")
class TestExcelExport(UrlInventoryTestBase):
    def test_export(self):
        from openpyxl import load_workbook

        path = self.tmp / "url_inventory.xlsx"
        result = ui.export_inventory_to_excel(self.inv, path)
        self.assertEqual(Path(result), path)
        self.assertTrue(path.exists())

        wb = load_workbook(path)
        self.assertEqual(
            wb.sheetnames, ["サマリ", "Webhook", "JS呼び出し", "外部参照"]
        )

        ws = wb["Webhook"]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers[:5], ["アプリID", "アプリ名", "Webhook名", "有効", "URL"])
        self.assertEqual(ws.max_row, 1 + 2)

        ws = wb["JS呼び出し"]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(
            headers,
            [
                "アプリID",
                "アプリ名",
                "ファイル",
                "行",
                "種別",
                "メソッド",
                "URL（式）",
                "解決したURL",
                "ホスト",
                "状態",
                "変更後URL",
                "コード",
            ],
        )
        self.assertEqual(ws.max_row, 1 + 10)

        ws = wb["外部参照"]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers[:4], ["アプリID", "アプリ名", "種別", "画面"])
        self.assertEqual(ws.max_row, 1 + 3)

        summary = wb["サマリ"]
        values = [
            str(row[0]) for row in summary.iter_rows(min_col=1, max_col=1, values_only=True)
        ]
        self.assertIn("メタ情報", values)
        self.assertIn("カテゴリ別件数", values)
        self.assertIn("ホスト別", values)

    def test_status_fill(self):
        from openpyxl import load_workbook

        path = self.tmp / "url_inventory_fill.xlsx"
        ui.export_inventory_to_excel(self.inv, path)
        wb = load_workbook(path)
        ws = wb["Webhook"]
        status_col = [cell.value for cell in ws[1]].index("状態") + 1
        colors = set()
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=status_col)
            if cell.value == ui.STATUS_TODO:
                colors.add(cell.fill.start_color.rgb)
        self.assertTrue(any("FFC7CE" in str(color) for color in colors))


if __name__ == "__main__":
    unittest.main()
