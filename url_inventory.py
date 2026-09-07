#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""取得済み output から「外部 URL を持つ設定」を洗い出す（ドメイン名変更の棚卸し）。

対象は 3 カテゴリ。

1. Webhook      : 各アプリの Webhook 設定（保存済みファイルから）
2. JS呼び出し   : カスタマイズ JavaScript 内の HTTP 呼び出しと URL 文字列
3. 外部参照     : カスタマイズ設定の URL 指定 JS/CSS、カスタムビュー HTML 内の URL

.kintone.env の任意設定 `domain_map`（変更前ホスト → 変更後ホスト）があれば、
各行に「変更後 URL」と「状態」（要変更 / 変更済 / 対象外）を付ける。
変更前は「要変更」の洗い出しに、変更後は「要変更 0 件」の確認に使う。
"""

from __future__ import annotations

import bisect
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple

from console import BOLD, CYAN, DIM, GREEN, YELLOW, c, display_width
from inspect_downloaded import (
    extract_webhook_rows,
    find_webhook_file,
    is_stale_webhook_file,
    load_json_or_yaml,
)

# ---------------------------------------------------------------------------
# 定数
# ---------------------------------------------------------------------------

CAT_WEBHOOK = "Webhook"
CAT_JS = "JS呼び出し"
CAT_REF = "外部参照"

CATEGORIES = [CAT_WEBHOOK, CAT_JS, CAT_REF]

STATUS_TODO = "要変更"
STATUS_DONE = "変更済"
STATUS_NONE = "対象外"

STATUS_ORDER = {STATUS_TODO: 0, STATUS_DONE: 1, STATUS_NONE: 2}

HOST_DYNAMIC = "(動的)"
HOST_KINTONE = "(kintone)"
METHOD_UNKNOWN = "(変数)"
METHOD_NONE = "-"

# XML 名前空間など、ドメイン変更とは無関係のホスト
IGNORED_HOSTS = {
    "www.w3.org",
    "schemas.microsoft.com",
    "www.apache.org",
    "opensource.org",
}

# 1 行がこれを超える行があるファイルはミニファイ済みとみなして丸ごとスキップする
MINIFIED_LINE_LEN = 2000
MINIFIED_REASON = "ミニファイ済み（1 行が 2000 文字超）"

# 引数の切り出しはこの文字数で打ち切る
ARG_SCAN_LIMIT = 3000
# 式として記録する URL の最大長
MAX_EXPR_LEN = 300
MAX_CODE_LEN = 200


# ---------------------------------------------------------------------------
# データ構造
# ---------------------------------------------------------------------------


@dataclass
class UrlRow:
    app_id: str
    app_name: str
    category: str        # "Webhook" | "JS呼び出し" | "外部参照"
    kind: str            # 種別（fetch / kintone.proxy / カスタマイズJS ...）
    method: str          # "GET"/"POST"/... 大文字。不明・変数なら "(変数)"、該当なしは "-"
    url: str             # URL 文字列、または呼び出しの URL 引数の式そのもの
    resolved_url: str    # 式を変数定義から解決できた URL。できなければ ""
    host: str            # 抽出したホスト（小文字）。不明 ""、動的 "(動的)"、kintone "(kintone)"
    file: str            # app_dir からの相対パス（Webhook は ""）
    line: int            # 1 始まり（Webhook は 0）
    code: str            # 該当行を strip して 200 文字まで（Webhook はイベント一覧）
    extra: Dict[str, str] = field(default_factory=dict)
    to_be_url: str = ""  # domain_map 適用後の URL（対象外なら ""）
    status: str = ""     # "要変更" / "変更済" / "対象外"（domain_map が空なら ""）

    @property
    def effective_url(self) -> str:
        """置換対象の URL。解決できていればそちらを使う。"""
        return self.resolved_url or self.url


@dataclass
class Inventory:
    rows: List[UrlRow]
    apps: List[Tuple[str, str]]                 # (app_id, app_name) スキャン対象
    missing_webhooks: List[str]                 # Webhook 未取得のアプリID
    skipped_files: List[Tuple[str, str, str]]   # (app_id, file, 理由)
    domain_map: Dict[str, str]
    generated_at: datetime


# ---------------------------------------------------------------------------
# 小さなヘルパー
# ---------------------------------------------------------------------------


def _app_name_from_dir(app_id: str, app_dir: Path) -> str:
    """フォルダ名 [アプリID]_[アプリ名] の `_` 以降をアプリ名とする。"""
    name = app_dir.name
    prefix = f"{app_id}_"
    if name.startswith(prefix):
        return name[len(prefix):]
    parts = name.split("_", 1)
    return parts[1] if len(parts) == 2 else name


def _app_sort_key(app_id: str) -> Tuple[int, int, str]:
    text = str(app_id or "")
    if text.isdigit():
        return (0, int(text), "")
    return (1, 0, text)


def _row_sort_key(row: "UrlRow") -> Tuple[Any, ...]:
    return (_app_sort_key(row.app_id), row.file, row.line, row.kind, row.url)


def _shorten(text: str, limit: int) -> str:
    text = (text or "").strip()
    if len(text) > limit:
        return text[: limit - 3] + "..."
    return text


def _collapse(expr: str) -> str:
    return _shorten(re.sub(r"\s+", " ", (expr or "").strip()), MAX_EXPR_LEN)


def _line_starts(text: str) -> List[int]:
    starts = [0]
    for match in re.finditer(r"\n", text):
        starts.append(match.end())
    return starts


def _line_no(index: int, starts: Sequence[int]) -> int:
    return bisect.bisect_right(starts, index)


def _is_comment_line(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("//") or stripped.startswith("*") or stripped.startswith("/*")


_LITERAL_RE = re.compile(r"^(['\"`])(.*)\1$", re.S)


def _string_literal(expr: str) -> Optional[str]:
    """文字列リテラルなら中身を返す。テンプレート内に ${} があればリテラル扱いしない。"""
    text = (expr or "").strip()
    match = _LITERAL_RE.match(text)
    if not match:
        return None
    inner = match.group(2)
    quote = match.group(1)
    if quote in ("'", '"') and quote in inner.replace("\\" + quote, ""):
        return None
    if quote == "`" and "${" in inner:
        return None
    return inner


def _extract_host(text: str) -> str:
    """URL・式からホストを取り出して小文字で返す。動的なら "(動的)"。"""
    if not text:
        return ""
    match = re.search(r"(?:https?:)?//([^/'\"`\s:?#]*)", text)
    if not match:
        return ""
    host = match.group(1).strip()
    if not host:
        # 'https://' + host のように連結しているもの
        return HOST_DYNAMIC
    if "${" in host or "+" in host or "}" in host:
        return HOST_DYNAMIC
    host = host.split("@")[-1]
    if not host:
        return HOST_DYNAMIC
    return host.lower()


def _replace_host(url: str, old_host: str, new_host: str) -> str:
    if not url or not old_host or not new_host:
        return ""
    return re.sub(re.escape(old_host), new_host, url, count=1, flags=re.IGNORECASE)


# ---------------------------------------------------------------------------
# domain_map
# ---------------------------------------------------------------------------


def normalize_domain_map(raw: Any) -> Dict[str, str]:
    """.kintone.env の domain_map を {小文字ホスト: 小文字ホスト} に正規化する。

    'https://' や末尾 '/' が付いていても取り除く。None / 非 dict は {}。
    """
    if not isinstance(raw, dict):
        return {}
    result: Dict[str, str] = {}
    for key, value in raw.items():
        before = _normalize_host_text(key)
        after = _normalize_host_text(value)
        if before and after:
            result[before] = after
    return result


def _normalize_host_text(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    text = re.sub(r"^[A-Za-z][A-Za-z0-9+.\-]*:", "", text)
    text = text.lstrip("/")
    text = text.split("/", 1)[0]
    text = text.split("?", 1)[0]
    text = text.split("#", 1)[0]
    text = text.split("@")[-1]
    text = text.split(":", 1)[0]
    return text.strip().strip(".").lower()


def _to_be_host(host: str, domain_map: Dict[str, str]) -> Tuple[str, str]:
    """(状態, 変更後ホスト) を返す。domain_map が空なら ("", "")。"""
    if not domain_map:
        return ("", "")
    if not host or host.startswith("("):
        return (STATUS_NONE, "")
    for before in sorted(domain_map, key=len, reverse=True):
        after = domain_map[before]
        if host == before:
            return (STATUS_TODO, after)
        if host.endswith("." + before):
            return (STATUS_TODO, host[: -len(before)] + after)
    for after in domain_map.values():
        if host == after or host.endswith("." + after):
            return (STATUS_DONE, "")
    return (STATUS_NONE, "")


def _apply_domain_map(row: UrlRow, domain_map: Dict[str, str]) -> None:
    status, new_host = _to_be_host(row.host, domain_map)
    row.status = status
    if status == STATUS_TODO and new_host:
        row.to_be_url = _replace_host(row.effective_url, row.host, new_host)
    else:
        row.to_be_url = ""


# ---------------------------------------------------------------------------
# JavaScript の走査
# ---------------------------------------------------------------------------


_DEF_RE = re.compile(
    r"(?:const|let|var)\s+([A-Za-z_$][\w$]*)\s*=\s*['\"`]((?:https?:)?//[^'\"`\n]*)['\"`]"
)
_PROP_RE = re.compile(
    r"(?<![\w$])([A-Za-z_$][\w$]*)\s*:\s*['\"`]((?:https?:)?//[^'\"`\n]*)['\"`]"
)

_STR_URL_RE = re.compile(
    r"'([^'\n]*https?://[^'\n]*)'"
    r"|\"([^\"\n]*https?://[^\"\n]*)\""
    r"|`([^`\n]*https?://[^`\n]*)`"
)

_XHR_RECEIVER_HINTS = ("xhr", "req", "request", "http")
_XHR_RECEIVER_DENY = {"window", "document", "self", "top", "parent", "opener", "dialog"}

_CALL_PATTERNS: List[Tuple[str, "re.Pattern[str]"]] = [
    ("fetch", re.compile(r"\bfetch\s*\(")),
    ("kintone.proxy", re.compile(r"\bkintone\s*\.\s*proxy\s*\(")),
    ("kintone.api", re.compile(r"\bkintone\s*\.\s*api\s*\(")),
    ("jquery.ajax", re.compile(r"(?:\$|jQuery)\s*\.\s*ajax\s*\(")),
    ("jquery.short", re.compile(r"(?:\$|jQuery)\s*\.\s*(get|post|getJSON)\s*\(")),
    ("axios.method", re.compile(r"\baxios\s*\.\s*(get|post|put|delete|patch|head)\s*\(")),
    ("axios", re.compile(r"\baxios\s*\(")),
    ("sendBeacon", re.compile(r"\bnavigator\s*\.\s*sendBeacon\s*\(")),
    ("websocket", re.compile(r"\bnew\s+WebSocket\s*\(")),
    ("window.open", re.compile(r"\bwindow\s*\.\s*open\s*\(")),
    ("location.call", re.compile(r"\blocation\s*\.\s*(assign|replace)\s*\(")),
    ("location.href", re.compile(r"\blocation\s*\.\s*href\s*=\s*")),
    ("xhr", re.compile(r"([A-Za-z_$][\w$]*)\s*\.\s*open\s*\(")),
]


def _split_call_args(text: str, open_idx: int, limit: int = ARG_SCAN_LIMIT) -> List[str]:
    """text[open_idx] == '(' として、対応する ')' までをトップレベルのカンマで分割する。

    文字列（' " `）とネストした括弧、テンプレートリテラルの ${} を考慮する。
    """
    end = min(len(text), open_idx + 1 + limit)
    i = open_idx + 1
    start = i
    depth = 1
    quote: Optional[str] = None
    tmpl: List[Tuple[str, int]] = []
    args: List[str] = []
    closed = False
    while i < end:
        ch = text[i]
        if quote is not None:
            if ch == "\\":
                i += 2
                continue
            if quote == "`" and ch == "$" and i + 1 < end and text[i + 1] == "{":
                tmpl.append((quote, depth))
                quote = None
                depth += 1
                i += 2
                continue
            if ch == quote:
                quote = None
            i += 1
            continue
        if ch in "'\"`":
            quote = ch
            i += 1
            continue
        if ch in "([{":
            depth += 1
            i += 1
            continue
        if ch in ")]}":
            depth -= 1
            if depth == 0:
                args.append(text[start:i])
                closed = True
                break
            if ch == "}" and tmpl and tmpl[-1][1] == depth:
                quote = tmpl.pop()[0]
            i += 1
            continue
        if ch == "," and depth == 1:
            args.append(text[start:i])
            start = i + 1
            i += 1
            continue
        i += 1
    if not closed and start < end:
        args.append(text[start:end])
    cleaned = [arg.strip() for arg in args]
    if len(cleaned) == 1 and not cleaned[0]:
        return []
    return cleaned


def _object_value(text: str, key: str) -> str:
    """オブジェクトリテラル風のテキストから key の値の式を取り出す。"""
    if not text:
        return ""
    pattern = re.compile(r"(?<![\w$.])['\"]?" + re.escape(key) + r"['\"]?\s*:\s*")
    match = pattern.search(text)
    if not match:
        return ""
    i = match.end()
    depth = 0
    quote: Optional[str] = None
    while i < len(text):
        ch = text[i]
        if quote is not None:
            if ch == "\\":
                i += 2
                continue
            if ch == quote:
                quote = None
            i += 1
            continue
        if ch in "'\"`":
            quote = ch
            i += 1
            continue
        if ch in "([{":
            depth += 1
        elif ch in ")]}":
            if depth == 0:
                break
            depth -= 1
        elif ch == "," and depth == 0:
            break
        i += 1
    return text[match.end():i].strip()


def _method_from_options(options: str, keys: Sequence[str], default: str) -> str:
    """オプションオブジェクトから method / type を取り出す。"""
    if not options:
        return default
    for key in keys:
        value = _object_value(options, key)
        if not value:
            continue
        literal = _string_literal(value)
        if literal:
            return literal.strip().upper()
        return METHOD_UNKNOWN
    return default


def _method_from_arg(arg: str) -> str:
    if not arg:
        return METHOD_UNKNOWN
    literal = _string_literal(arg)
    if literal:
        return literal.strip().upper() or METHOD_UNKNOWN
    return METHOD_UNKNOWN


def _collect_definitions(text: str) -> Dict[str, str]:
    """同一ファイル内の `const X = 'https://...'` / `x: 'https://...'` を集める。"""
    defs: Dict[str, str] = {}
    for match in _DEF_RE.finditer(text):
        defs.setdefault(match.group(1), match.group(2))
    for match in _PROP_RE.finditer(text):
        defs.setdefault(match.group(1), match.group(2))
    return defs


def _resolve_expr(expr: str, defs: Dict[str, str]) -> str:
    """識別子 / IDENT + '/path' / `${IDENT}/path` を 1 段階だけ解決する。"""
    text = (expr or "").strip()
    if not text or not defs:
        return ""
    match = re.fullmatch(r"([A-Za-z_$][\w$]*)", text)
    if match:
        return defs.get(match.group(1), "")
    match = re.fullmatch(r"([A-Za-z_$][\w$]*)\s*\+\s*['\"`]([^'\"`]*)['\"`]", text)
    if match:
        base = defs.get(match.group(1), "")
        return base + match.group(2) if base else ""
    match = re.fullmatch(r"`\$\{\s*([A-Za-z_$][\w$]*)\s*\}([^`]*)`", text)
    if match:
        base = defs.get(match.group(1), "")
        return base + match.group(2) if base else ""
    return ""


def _url_parts(expr: str, defs: Dict[str, str]) -> Tuple[str, str, str]:
    """URL 引数の式から (url, resolved_url, host) を作る。"""
    literal = _string_literal(expr)
    if literal is not None:
        url = literal.strip()
        return (url, "", _extract_host(url))
    url = _collapse(expr)
    resolved = _resolve_expr(expr, defs)
    host = _extract_host(resolved) if resolved else _extract_host(url)
    return (url, resolved, host)


def _scan_js_text(text: str, rel_path: str) -> List[Dict[str, Any]]:
    """JavaScript のテキストから呼び出しと URL 文字列を拾う。"""
    lines = text.splitlines()
    starts = _line_starts(text)
    defs = _collect_definitions(text)
    found: List[Dict[str, Any]] = []
    call_lines: Set[int] = set()

    matches: List[Tuple[int, str, Any]] = []
    for tag, pattern in _CALL_PATTERNS:
        for match in pattern.finditer(text):
            matches.append((match.start(), tag, match))
    matches.sort(key=lambda item: (item[0], item[1]))

    for start, tag, match in matches:
        line_no = _line_no(start, starts)
        line_text = lines[line_no - 1] if 0 < line_no <= len(lines) else ""
        if _is_comment_line(line_text):
            continue

        kind = ""
        url_expr = ""
        method = METHOD_NONE

        if tag == "fetch":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "fetch"
            url_expr = args[0]
            method = _method_from_options(args[1] if len(args) > 1 else "", ("method",), "GET")
        elif tag == "kintone.proxy":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "kintone.proxy"
            url_expr = args[0]
            method = _method_from_arg(args[1]) if len(args) > 1 else METHOD_UNKNOWN
        elif tag == "kintone.api":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "kintone.api"
            url_expr = args[0]
            method = _method_from_arg(args[1]) if len(args) > 1 else METHOD_UNKNOWN
        elif tag == "jquery.ajax":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "jQuery.ajax"
            first_literal = _string_literal(args[0])
            if first_literal is not None:
                url_expr = args[0]
                options = args[1] if len(args) > 1 else ""
            else:
                options = args[0]
                url_expr = _object_value(options, "url")
            method = _method_from_options(options, ("type", "method"), "GET")
            if not url_expr:
                continue
        elif tag == "jquery.short":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            name = match.group(1)
            kind = "jQuery." + name
            url_expr = args[0]
            method = "POST" if name == "post" else "GET"
        elif tag == "axios.method":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "axios"
            url_expr = args[0]
            method = match.group(1).upper()
        elif tag == "axios":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "axios"
            first_literal = _string_literal(args[0])
            if first_literal is not None:
                url_expr = args[0]
                options = args[1] if len(args) > 1 else ""
            else:
                options = args[0]
                url_expr = _object_value(options, "url")
            method = _method_from_options(options, ("method",), "GET")
            if not url_expr:
                continue
        elif tag == "sendBeacon":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "sendBeacon"
            url_expr = args[0]
            method = "POST"
        elif tag == "websocket":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "WebSocket"
            url_expr = args[0]
            method = METHOD_NONE
        elif tag == "window.open":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "window.open"
            url_expr = args[0]
            method = METHOD_NONE
        elif tag == "location.call":
            args = _split_call_args(text, match.end() - 1)
            if not args:
                continue
            kind = "location"
            url_expr = args[0]
            method = METHOD_NONE
        elif tag == "location.href":
            rest = text[match.end():]
            rest = rest.split(";", 1)[0].split("\n", 1)[0]
            if not rest.strip():
                continue
            kind = "location"
            url_expr = rest.strip()
            method = METHOD_NONE
        elif tag == "xhr":
            receiver = (match.group(1) or "").lower()
            if receiver in _XHR_RECEIVER_DENY:
                continue
            near = text[max(0, start - 300):start]
            if "XMLHttpRequest" not in near and not any(
                hint in receiver for hint in _XHR_RECEIVER_HINTS
            ):
                continue
            args = _split_call_args(text, match.end() - 1)
            if len(args) < 2:
                continue
            kind = "XMLHttpRequest"
            method = _method_from_arg(args[0])
            url_expr = args[1]
        else:
            continue

        url, resolved, host = _url_parts(url_expr, defs)
        if kind == "kintone.api" and not host:
            host = HOST_KINTONE
        if host in IGNORED_HOSTS:
            continue
        if not url:
            continue
        call_lines.add(line_no)
        found.append(
            {
                "kind": kind,
                "method": method or METHOD_NONE,
                "url": url,
                "resolved_url": resolved,
                "host": host,
                "file": rel_path,
                "line": line_no,
                "code": _shorten(line_text, MAX_CODE_LEN),
            }
        )

    # 呼び出しに該当しない行の URL 文字列
    for index, line_text in enumerate(lines, start=1):
        if index in call_lines or _is_comment_line(line_text):
            continue
        if "http" not in line_text:
            continue
        for match in _STR_URL_RE.finditer(line_text):
            url = (match.group(1) or match.group(2) or match.group(3) or "").strip()
            if not url:
                continue
            host = _extract_host(url)
            if host in IGNORED_HOSTS:
                continue
            found.append(
                {
                    "kind": "URL文字列",
                    "method": METHOD_NONE,
                    "url": _shorten(url, MAX_EXPR_LEN),
                    "resolved_url": "",
                    "host": host,
                    "file": rel_path,
                    "line": index,
                    "code": _shorten(line_text, MAX_CODE_LEN),
                }
            )

    # 同一 (file, line, url, kind) の重複はまとめる
    unique: List[Dict[str, Any]] = []
    seen: Set[Tuple[str, int, str, str]] = set()
    for item in sorted(found, key=lambda x: (x["line"], x["kind"], x["url"])):
        key = (item["file"], item["line"], item["url"], item["kind"])
        if key in seen:
            continue
        seen.add(key)
        unique.append(item)
    return unique


def scan_javascript(app_dir: Path) -> Tuple[List[Dict[str, Any]], List[Tuple[str, str]]]:
    """javascript/ 配下の *.js を走査して (検出, スキップしたファイル) を返す。"""
    js_dir = app_dir / "javascript"
    hits: List[Dict[str, Any]] = []
    skipped: List[Tuple[str, str]] = []
    if not js_dir.is_dir():
        return (hits, skipped)
    for path in sorted(js_dir.rglob("*.js")):
        if not path.is_file():
            continue
        rel = path.relative_to(app_dir).as_posix()
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            skipped.append((rel, f"読み込みできません（{exc}）"))
            continue
        if any(len(line) > MINIFIED_LINE_LEN for line in text.splitlines()):
            skipped.append((rel, MINIFIED_REASON))
            continue
        hits.extend(_scan_js_text(text, rel))
    return (hits, skipped)


# ---------------------------------------------------------------------------
# カスタマイズ設定 / ビュー
# ---------------------------------------------------------------------------


def _find_config_file(app_dir: Path, app_id: str, name: str) -> Optional[Path]:
    candidates = [
        app_dir / f"{app_id}_{name}.yaml",
        app_dir / f"{app_id}_{name}.yml",
        app_dir / f"{app_id}_{name}.json",
        app_dir / "json" / f"{app_id}_{name}.json",
        app_dir / f"{name}.yaml",
        app_dir / f"{name}.json",
    ]
    for path in candidates:
        if path.is_file():
            return path
    return None


def _text_lines(path: Path) -> List[str]:
    try:
        return path.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError:
        return []


def _find_line(lines: Sequence[str], needle: str, used: Set[int]) -> int:
    """テキスト中で needle を含む未使用の行番号（1 始まり）。無ければ 0。"""
    if not needle:
        return 0
    for index, line in enumerate(lines, start=1):
        if index in used:
            continue
        if needle in line:
            used.add(index)
            return index
    for index, line in enumerate(lines, start=1):
        if needle in line:
            return index
    return 0


_HTML_URL_RE = re.compile(r"https?://[^\s\"'<>)\\]+")


def scan_customize(app_dir: Path, app_id: str) -> List[Dict[str, Any]]:
    """カスタマイズ設定の type: URL の JS/CSS を拾う。"""
    path = _find_config_file(app_dir, app_id, "customize")
    if path is None:
        return []
    try:
        data = load_json_or_yaml(path)
    except Exception:
        return []
    if not isinstance(data, dict):
        return []
    rel = path.relative_to(app_dir).as_posix()
    lines = _text_lines(path)
    used: Set[int] = set()
    rows: List[Dict[str, Any]] = []
    for platform in ("desktop", "mobile"):
        section = data.get(platform)
        if not isinstance(section, dict):
            continue
        for file_type in ("js", "css"):
            items = section.get(file_type)
            if isinstance(items, dict):
                items = [items]
            if not isinstance(items, list):
                continue
            for item in items:
                if not isinstance(item, dict):
                    continue
                if str(item.get("type") or "").upper() != "URL":
                    continue
                url = str(item.get("url") or "").strip()
                if not url:
                    continue
                host = _extract_host(url)
                if host in IGNORED_HOSTS:
                    continue
                line_no = _find_line(lines, url, used)
                code = lines[line_no - 1].strip() if line_no else url
                rows.append(
                    {
                        "kind": "カスタマイズJS" if file_type == "js" else "カスタマイズCSS",
                        "method": METHOD_NONE,
                        "url": url,
                        "resolved_url": "",
                        "host": host,
                        "file": rel,
                        "line": line_no,
                        "code": _shorten(code, MAX_CODE_LEN),
                        "extra": {"platform": platform, "file_type": file_type},
                    }
                )
    return rows


def _iter_views(data: Any):
    if isinstance(data, dict):
        views = data.get("views", data)
    else:
        views = data
    if isinstance(views, dict):
        for name, view in views.items():
            if isinstance(view, dict):
                yield (str(view.get("name") or name), view)
    elif isinstance(views, list):
        for view in views:
            if isinstance(view, dict):
                yield (str(view.get("name") or view.get("id") or ""), view)


def scan_views(app_dir: Path, app_id: str) -> List[Dict[str, Any]]:
    """カスタムビューの HTML 内の URL を拾う。"""
    path = _find_config_file(app_dir, app_id, "views")
    if path is None:
        return []
    try:
        data = load_json_or_yaml(path)
    except Exception:
        return []
    rel = path.relative_to(app_dir).as_posix()
    lines = _text_lines(path)
    used: Set[int] = set()
    rows: List[Dict[str, Any]] = []
    seen: Set[Tuple[str, str]] = set()
    for view_name, view in _iter_views(data):
        if str(view.get("type") or "").upper() != "CUSTOM":
            continue
        html = view.get("html")
        if not isinstance(html, str) or not html:
            continue
        html_lines = html.splitlines() or [html]
        for offset, html_line in enumerate(html_lines, start=1):
            for match in _HTML_URL_RE.finditer(html_line):
                url = match.group(0).rstrip(".,;:")
                host = _extract_host(url)
                if host in IGNORED_HOSTS:
                    continue
                key = (view_name, url)
                if key in seen:
                    continue
                seen.add(key)
                line_no = _find_line(lines, url, used) or offset
                code = lines[line_no - 1].strip() if 0 < line_no <= len(lines) else html_line.strip()
                rows.append(
                    {
                        "kind": "ビューHTML",
                        "method": METHOD_NONE,
                        "url": url,
                        "resolved_url": "",
                        "host": host,
                        "file": rel,
                        "line": line_no,
                        "code": _shorten(code, MAX_CODE_LEN),
                        "extra": {"platform": view_name, "file_type": "html"},
                    }
                )
    return rows


# ---------------------------------------------------------------------------
# Webhook
# ---------------------------------------------------------------------------


def scan_webhooks(app_dir: Path, app_id: str) -> Tuple[Optional[List[Dict[str, Any]]], str]:
    """(行, 理由) を返す。未取得のときは (None, 理由)。"""
    path = find_webhook_file(app_dir, app_id)
    if path is None:
        return (None, "Webhook ファイルがありません")
    try:
        data = load_json_or_yaml(path)
    except Exception as exc:
        return (None, f"Webhook ファイルを読めません（{exc}）")
    if is_stale_webhook_file(data):
        return (None, "Webhook ファイルが旧方式（再取得が必要）")
    return (extract_webhook_rows(data), "")


def _enabled_text(value: Any) -> str:
    if value is True:
        return "有効"
    if value is False:
        return "無効"
    text = str(value or "").strip()
    if text.lower() in ("true", "1"):
        return "有効"
    if text.lower() in ("false", "0"):
        return "無効"
    return text or "-"


# ---------------------------------------------------------------------------
# 収集
# ---------------------------------------------------------------------------


def build_inventory(
    targets: List[Tuple[str, Path]],
    domain_map: Optional[Dict[str, str]] = None,
) -> Inventory:
    """取得済みフォルダ一覧から URL 一覧を作る。

    targets は inspect_downloaded.list_app_output_dirs() と同じ [(app_id, app_dir)]。
    """
    dmap = normalize_domain_map(domain_map)
    rows: List[UrlRow] = []
    apps: List[Tuple[str, str]] = []
    missing_webhooks: List[str] = []
    skipped_files: List[Tuple[str, str, str]] = []

    for app_id, app_dir in targets:
        app_dir = Path(app_dir)
        app_name = _app_name_from_dir(app_id, app_dir)
        apps.append((app_id, app_name))

        # 1. Webhook
        webhook_rows, _reason = scan_webhooks(app_dir, app_id)
        if webhook_rows is None:
            missing_webhooks.append(app_id)
        else:
            for item in webhook_rows:
                url = str(item.get("url") or "").strip()
                rows.append(
                    UrlRow(
                        app_id=app_id,
                        app_name=app_name,
                        category=CAT_WEBHOOK,
                        kind="Webhook",
                        method=METHOD_NONE,
                        url=url,
                        resolved_url="",
                        host=_extract_host(url),
                        file="",
                        line=0,
                        code=str(item.get("events") or ""),
                        extra={
                            "id": str(item.get("id") or ""),
                            "name": str(item.get("name") or ""),
                            "enabled": _enabled_text(item.get("enabled")),
                            "events": str(item.get("events") or ""),
                            "modifier": str(item.get("modifier") or ""),
                            "modified_at": str(item.get("modified_at") or ""),
                        },
                    )
                )

        # 2. JS 呼び出し
        js_hits, skipped = scan_javascript(app_dir)
        for rel, reason in skipped:
            skipped_files.append((app_id, rel, reason))
        for item in js_hits:
            rows.append(
                UrlRow(
                    app_id=app_id,
                    app_name=app_name,
                    category=CAT_JS,
                    kind=item["kind"],
                    method=item["method"],
                    url=item["url"],
                    resolved_url=item["resolved_url"],
                    host=item["host"],
                    file=item["file"],
                    line=item["line"],
                    code=item["code"],
                )
            )

        # 3. 外部参照
        for item in list(scan_customize(app_dir, app_id)) + list(scan_views(app_dir, app_id)):
            rows.append(
                UrlRow(
                    app_id=app_id,
                    app_name=app_name,
                    category=CAT_REF,
                    kind=item["kind"],
                    method=item["method"],
                    url=item["url"],
                    resolved_url=item["resolved_url"],
                    host=item["host"],
                    file=item["file"],
                    line=item["line"],
                    code=item["code"],
                    extra=dict(item.get("extra") or {}),
                )
            )

    for row in rows:
        _apply_domain_map(row, dmap)

    rows.sort(key=_row_sort_key)
    return Inventory(
        rows=rows,
        apps=apps,
        missing_webhooks=missing_webhooks,
        skipped_files=skipped_files,
        domain_map=dmap,
        generated_at=datetime.now(),
    )


# ---------------------------------------------------------------------------
# 集計
# ---------------------------------------------------------------------------


def summarize_hosts(inv: Inventory) -> List[Dict[str, object]]:
    """ホスト別の集計。要変更 → 変更済 → 対象外 の順、同じ状態内は合計の降順。"""
    grouped: Dict[str, Dict[str, Any]] = {}
    for row in inv.rows:
        entry = grouped.setdefault(
            row.host,
            {
                "host": row.host,
                "webhook": 0,
                "js": 0,
                "ref": 0,
                "total": 0,
                "apps": set(),
                "status": row.status,
                "to_be_host": "",
            },
        )
        if row.category == CAT_WEBHOOK:
            entry["webhook"] += 1
        elif row.category == CAT_JS:
            entry["js"] += 1
        else:
            entry["ref"] += 1
        entry["total"] += 1
        entry["apps"].add(row.app_id)

    result: List[Dict[str, object]] = []
    for host, entry in grouped.items():
        status, to_be_host = _to_be_host(host, inv.domain_map)
        entry["status"] = status
        entry["to_be_host"] = to_be_host
        entry["apps"] = sorted(entry["apps"], key=_app_sort_key)
        result.append(entry)
    result.sort(
        key=lambda item: (
            STATUS_ORDER.get(str(item["status"]), 3),
            -int(item["total"]),
            str(item["host"]),
        )
    )
    return result


def count_by_category(inv: Inventory) -> Dict[str, int]:
    counts = {name: 0 for name in CATEGORIES}
    for row in inv.rows:
        counts[row.category] = counts.get(row.category, 0) + 1
    counts["合計"] = len(inv.rows)
    return counts


def count_by_status(inv: Inventory) -> Dict[str, int]:
    counts = {STATUS_TODO: 0, STATUS_DONE: 0, STATUS_NONE: 0}
    for row in inv.rows:
        if row.status:
            counts[row.status] = counts.get(row.status, 0) + 1
    return counts


def rows_to_change(inv: Inventory) -> List[UrlRow]:
    return [row for row in inv.rows if row.status == STATUS_TODO]


# ---------------------------------------------------------------------------
# 端末表示
# ---------------------------------------------------------------------------


def _app_label(row: UrlRow) -> str:
    return f"{row.app_id}_{row.app_name}" if row.app_name else str(row.app_id)


def format_inventory_summary(inv: Inventory) -> List[str]:
    """端末表示用の行リスト（色付き）。"""
    lines: List[str] = []
    lines.append(c("  URL 一覧（外部URLを持つ設定の棚卸し）", BOLD, CYAN))
    lines.append(
        c(f"  生成日時: {inv.generated_at.strftime('%Y-%m-%d %H:%M:%S')}", DIM)
    )
    app_ids = ", ".join(app_id for app_id, _ in inv.apps)
    lines.append(f"  対象アプリ: {len(inv.apps)} 件  ({app_ids or '-'})")

    counts = count_by_category(inv)
    lines.append(
        c(
            "  検出: 合計 {total} 件"
            "  （{w}: {wc} / {j}: {jc} / {r}: {rc}）".format(
                total=counts["合計"],
                w=CAT_WEBHOOK,
                wc=counts[CAT_WEBHOOK],
                j=CAT_JS,
                jc=counts[CAT_JS],
                r=CAT_REF,
                rc=counts[CAT_REF],
            ),
            BOLD,
        )
    )

    if inv.domain_map:
        pairs = ", ".join(f"{k} -> {v}" for k, v in sorted(inv.domain_map.items()))
        lines.append(f"  domain_map: {pairs}")
        status_counts = count_by_status(inv)
        todo = status_counts[STATUS_TODO]
        style = (BOLD, YELLOW) if todo else (BOLD, GREEN)
        lines.append(
            c(
                "  状態: {t} {tc} 件 / {d} {dc} 件 / {n} {nc} 件".format(
                    t=STATUS_TODO,
                    tc=todo,
                    d=STATUS_DONE,
                    dc=status_counts[STATUS_DONE],
                    n=STATUS_NONE,
                    nc=status_counts[STATUS_NONE],
                ),
                *style,
            )
        )
    else:
        lines.append(c("  domain_map: 未設定（状態の判定なし）", DIM))

    hosts = summarize_hosts(inv)
    if hosts:
        lines.append("")
        lines.append(c("  ホスト別:", BOLD))
        for entry in hosts:
            host = str(entry["host"]) or "(不明)"
            label = f"    ・{host}"
            status = str(entry["status"])
            if status == STATUS_TODO:
                label += c(f"  [{status} -> {entry['to_be_host']}]", BOLD, YELLOW)
            elif status == STATUS_DONE:
                label += c(f"  [{status}]", GREEN)
            elif status:
                label += c(f"  [{status}]", DIM)
            label += (
                f"  合計 {entry['total']}"
                f"（{CAT_WEBHOOK} {entry['webhook']} /"
                f" {CAT_JS} {entry['js']} /"
                f" {CAT_REF} {entry['ref']}）"
            )
            label += f"  アプリ: {', '.join(str(a) for a in entry['apps'])}"
            lines.append(label)

    todo_rows = rows_to_change(inv)
    if todo_rows:
        lines.append("")
        lines.append(c(f"  要変更 {len(todo_rows)} 件:", BOLD, YELLOW))
        for row in todo_rows[:30]:
            where = row.file or "-"
            if row.line:
                where = f"{where}:{row.line}"
            lines.append(
                f"    ・[{_app_label(row)}] {row.category}/{row.kind}  {where}"
            )
            lines.append(f"        {row.effective_url}  ->  {row.to_be_url}")
        if len(todo_rows) > 30:
            lines.append(f"    ・ほか {len(todo_rows) - 30} 件")
    elif inv.domain_map:
        lines.append("")
        lines.append(c("  要変更 0 件（domain_map の変更前ホストは残っていません）", BOLD, GREEN))

    if inv.missing_webhooks:
        lines.append("")
        lines.append(
            c(
                f"  △ Webhook 未取得: {len(inv.missing_webhooks)} アプリ"
                f"（{', '.join(inv.missing_webhooks)}）"
                "  … username / password を設定して app を実行すると取得されます",
                YELLOW,
            )
        )
    if inv.skipped_files:
        lines.append("")
        lines.append(c(f"  △ スキップしたファイル: {len(inv.skipped_files)} 件", YELLOW))
        for app_id, rel, reason in inv.skipped_files[:20]:
            lines.append(f"    ・[{app_id}] {rel}  {reason}")
        if len(inv.skipped_files) > 20:
            lines.append(f"    ・ほか {len(inv.skipped_files) - 20} 件")
    return lines


# ---------------------------------------------------------------------------
# Excel 出力
# ---------------------------------------------------------------------------


def _auto_widths(headers: Sequence[Any], rows: Sequence[Sequence[Any]]) -> List[int]:
    widths: List[int] = []
    for index, header in enumerate(headers):
        width = display_width(str(header))
        for row in rows:
            if index < len(row):
                value = row[index]
                width = max(width, display_width("" if value is None else str(value)))
        widths.append(max(8, min(80, width + 2)))
    return widths


def _webhook_table(inv: Inventory) -> Tuple[List[str], List[List[Any]]]:
    headers = [
        "アプリID",
        "アプリ名",
        "Webhook名",
        "有効",
        "URL",
        "ホスト",
        "状態",
        "変更後URL",
        "イベント",
        "ID",
        "更新者",
        "更新日時",
    ]
    rows = [
        [
            row.app_id,
            row.app_name,
            row.extra.get("name", ""),
            row.extra.get("enabled", ""),
            row.url,
            row.host,
            row.status,
            row.to_be_url,
            row.extra.get("events", ""),
            row.extra.get("id", ""),
            row.extra.get("modifier", ""),
            row.extra.get("modified_at", ""),
        ]
        for row in sorted(
            [r for r in inv.rows if r.category == CAT_WEBHOOK], key=_row_sort_key
        )
    ]
    return (headers, rows)


def _js_table(inv: Inventory) -> Tuple[List[str], List[List[Any]]]:
    headers = [
        "アプリID",
        "アプリ名",
        "ファイル",
        "行",
        "種別",
        "メソッド",
        "URL（式）",
        "解決したURL",
        "ホスト",
        "状態",
        "変更後URL",
        "コード",
    ]
    rows = [
        [
            row.app_id,
            row.app_name,
            row.file,
            row.line,
            row.kind,
            row.method,
            row.url,
            row.resolved_url,
            row.host,
            row.status,
            row.to_be_url,
            row.code,
        ]
        for row in sorted([r for r in inv.rows if r.category == CAT_JS], key=_row_sort_key)
    ]
    return (headers, rows)


def _ref_table(inv: Inventory) -> Tuple[List[str], List[List[Any]]]:
    headers = [
        "アプリID",
        "アプリ名",
        "種別",
        "画面",
        "ファイル",
        "行",
        "URL",
        "ホスト",
        "状態",
        "変更後URL",
    ]
    rows = [
        [
            row.app_id,
            row.app_name,
            row.kind,
            row.extra.get("platform", ""),
            row.file,
            row.line,
            row.url,
            row.host,
            row.status,
            row.to_be_url,
        ]
        for row in sorted([r for r in inv.rows if r.category == CAT_REF], key=_row_sort_key)
    ]
    return (headers, rows)


def export_inventory_to_excel(inv: Inventory, path: Path) -> Path:
    """URL 一覧を Excel に出力する（サマリ / Webhook / JS呼び出し / 外部参照）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    header_fill = PatternFill(start_color="E6F3FF", end_color="E6F3FF", fill_type="solid")
    section_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    todo_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    done_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_align = Alignment(vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    status_fills = {STATUS_TODO: todo_fill, STATUS_DONE: done_fill}

    def write_table(
        ws,
        headers: Sequence[Any],
        rows: Sequence[Sequence[Any]],
        start_row: int = 1,
        auto_filter: bool = True,
        freeze: bool = True,
    ) -> int:
        header_row = start_row
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin
        for offset, row in enumerate(rows):
            row_idx = header_row + 1 + offset
            for col, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.alignment = data_align
                cell.border = thin
                header = str(headers[col - 1]) if col <= len(headers) else ""
                if header == "状態":
                    fill = status_fills.get(str(value))
                    if fill is not None:
                        cell.fill = fill
        for col, width in enumerate(_auto_widths(headers, rows), 1):
            letter = get_column_letter(col)
            current = ws.column_dimensions[letter].width
            if current is None or current < width:
                ws.column_dimensions[letter].width = width
        last_row = header_row + len(rows)
        if auto_filter and rows:
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
        if freeze:
            ws.freeze_panes = f"A{header_row + 1}"
        ws.row_dimensions[header_row].height = 22
        return last_row

    def write_section(ws, row_idx: int, title: str, span: int) -> int:
        cell = ws.cell(row=row_idx, column=1, value=title)
        cell.font = header_font
        cell.fill = section_fill
        cell.border = thin
        for col in range(2, span + 1):
            other = ws.cell(row=row_idx, column=col)
            other.fill = section_fill
            other.border = thin
        if span > 1:
            ws.merge_cells(
                start_row=row_idx, start_column=1, end_row=row_idx, end_column=span
            )
        return row_idx + 1

    wb = Workbook()

    # --- サマリ ---
    ws = wb.active
    ws.title = "サマリ"
    app_ids = ", ".join(app_id for app_id, _ in inv.apps)
    domain_text = (
        ", ".join(f"{k} -> {v}" for k, v in sorted(inv.domain_map.items()))
        if inv.domain_map
        else "（未設定）"
    )
    meta_rows = [
        ["生成日時", inv.generated_at.strftime("%Y-%m-%d %H:%M:%S")],
        ["対象アプリ数", len(inv.apps)],
        ["対象アプリID", app_ids or "-"],
        [
            "Webhook 未取得アプリ",
            ", ".join(inv.missing_webhooks) if inv.missing_webhooks else "なし",
        ],
        ["スキップしたファイル数", len(inv.skipped_files)],
        ["domain_map", domain_text],
    ]
    row_idx = write_section(ws, 1, "メタ情報", 2)
    row_idx = write_table(ws, ["項目", "内容"], meta_rows, start_row=row_idx,
                          auto_filter=False, freeze=False) + 2

    counts = count_by_category(inv)
    status_counts = count_by_status(inv)
    count_rows: List[List[Any]] = [
        [CAT_WEBHOOK, counts[CAT_WEBHOOK]],
        [CAT_JS, counts[CAT_JS]],
        [CAT_REF, counts[CAT_REF]],
        ["合計", counts["合計"]],
        [STATUS_TODO, status_counts[STATUS_TODO]],
        [STATUS_DONE, status_counts[STATUS_DONE]],
        [STATUS_NONE, status_counts[STATUS_NONE]],
    ]
    row_idx = write_section(ws, row_idx, "カテゴリ別件数", 2)
    row_idx = write_table(ws, ["区分", "件数"], count_rows, start_row=row_idx,
                          auto_filter=False, freeze=False) + 2

    host_headers = [
        "ホスト",
        "状態",
        "変更後ホスト",
        CAT_WEBHOOK,
        CAT_JS,
        CAT_REF,
        "合計",
        "アプリ",
    ]
    host_rows = [
        [
            str(entry["host"]) or "(不明)",
            entry["status"],
            entry["to_be_host"],
            entry["webhook"],
            entry["js"],
            entry["ref"],
            entry["total"],
            ", ".join(str(a) for a in entry["apps"]),
        ]
        for entry in summarize_hosts(inv)
    ]
    row_idx = write_section(ws, row_idx, "ホスト別", len(host_headers))
    write_table(ws, host_headers, host_rows, start_row=row_idx,
                auto_filter=True, freeze=False)

    # --- 各カテゴリ ---
    for sheet_name, (headers, rows) in (
        (CAT_WEBHOOK, _webhook_table(inv)),
        (CAT_JS, _js_table(inv)),
        (CAT_REF, _ref_table(inv)),
    ):
        sheet = wb.create_sheet(sheet_name)
        write_table(sheet, headers, rows)

    wb.save(output_path)
    return output_path
